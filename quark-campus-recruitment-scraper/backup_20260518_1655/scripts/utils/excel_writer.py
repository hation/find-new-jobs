#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel生成器模块
将爬取的数据保存为Excel文件
"""

import os
import pandas as pd
import logging
from datetime import datetime
from typing import List, Dict, Any, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


class ExcelGenerator:
    """Excel生成器"""
    
    def __init__(self, config: Dict[str, Any]):
        """
        初始化Excel生成器
        
        Args:
            config: 配置字典
        """
        self.config = config
        self.output_config = config.get('output', {})
        self.output_dir = self.output_config.get('directory', './output')
        
        # 确保输出目录存在
        os.makedirs(self.output_dir, exist_ok=True)
        
    def generate_excel(self, data: List[Dict[str, Any]], 
                      filename: Optional[str] = None) -> str:
        """
        生成Excel文件
        
        Args:
            data: 数据列表
            filename: 自定义文件名（可选）
            
        Returns:
            生成的Excel文件路径
            
        Raises:
            ValueError: 数据为空时抛出
            IOError: 文件写入失败时抛出
        """
        if not data:
            raise ValueError("数据为空，无法生成Excel文件")
        
        # 生成文件名
        if filename is None:
            filename = self._generate_filename()
        
        filepath = os.path.join(self.output_dir, filename)
        
        logger.info(f"开始生成Excel文件: {filename}")
        logger.info(f"数据量: {len(data)} 条记录")
        
        try:
            # 创建DataFrame
            df = pd.DataFrame(data)
            
            # 保存Excel文件
            df.to_excel(filepath, index=False, engine='openpyxl')
            
            logger.info(f"Excel文件已保存: {filepath}")
            
            # 应用格式
            self._apply_excel_formatting(filepath, df)
            
            logger.info(f"Excel格式已应用: {filepath}")
            
            # 输出文件信息
            file_size = os.path.getsize(filepath)
            logger.info(f"文件大小: {file_size / 1024:.2f} KB")
            
            return filepath
            
        except Exception as e:
            logger.error(f"生成Excel文件失败: {str(e)}")
            raise IOError(f"Excel文件生成失败: {str(e)}")
    
    def _generate_filename(self) -> str:
        """
        生成文件名
        
        Returns:
            生成的文件名
        """
        prefix = self.output_config.get('filename_prefix', 'quark_positions')
        include_timestamp = self.output_config.get('include_timestamp', True)
        excel_format = self.output_config.get('excel_format', 'xlsx')
        
        if include_timestamp:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            return f"{prefix}_{timestamp}.{excel_format}"
        else:
            return f"{prefix}.{excel_format}"
    
    def _apply_excel_formatting(self, filepath: str, df: pd.DataFrame) -> None:
        """
        应用Excel格式
        
        Args:
            filepath: Excel文件路径
            df: 数据DataFrame
        """
        try:
            # 使用openpyxl加载工作簿
            from openpyxl import load_workbook
            
            wb = load_workbook(filepath)
            ws = wb.active
            
            # 设置列宽
            self._set_column_widths(ws, df)
            
            # 设置表头样式
            self._style_header(ws)
            
            # 设置数据样式
            self._style_data(ws, df)
            
            # 设置边框
            self._add_borders(ws, df)
            
            # 设置自动筛选
            self._add_auto_filter(ws, df)
            
            # 保存格式
            wb.save(filepath)
            
        except Exception as e:
            logger.warning(f"应用Excel格式失败（不影响数据）: {str(e)}")
    
    def _set_column_widths(self, ws, df: pd.DataFrame) -> None:
        """
        设置列宽
        
        Args:
            ws: 工作表对象
            df: 数据DataFrame
        """
        column_widths = self.output_config.get('column_widths', {})
        
        for i, column in enumerate(df.columns, 1):
            col_letter = get_column_letter(i)
            
            # 使用配置的列宽，或根据内容自动调整
            if column in column_widths:
                ws.column_dimensions[col_letter].width = column_widths[column]
            else:
                # 自动调整列宽
                max_length = 0
                column = str(column)
                
                # 计算表头长度
                max_length = max(max_length, len(column))
                
                # 计算数据长度
                for cell in ws[col_letter]:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                
                # 设置列宽（加一些缓冲）
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[col_letter].width = adjusted_width
    
    def _style_header(self, ws) -> None:
        """
        设置表头样式
        
        Args:
            ws: 工作表对象
        """
        header_font = Font(name='微软雅黑', size=12, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        # 应用样式到第一行
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
    
    def _style_data(self, ws, df: pd.DataFrame) -> None:
        """
        设置数据样式
        
        Args:
            ws: 工作表对象
            df: 数据DataFrame
        """
        data_font = Font(name='宋体', size=11)
        data_alignment = Alignment(vertical='top', wrap_text=True)
        
        # 为长文本列设置自动换行
        wrap_columns = ['职位描述', '职位要求']
        
        # 应用样式到数据行
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.font = data_font
                cell.alignment = data_alignment
                
                # 特定列设置自动换行
                col_index = cell.column - 1
                if col_index < len(df.columns):
                    column_name = df.columns[col_index]
                    if column_name in wrap_columns:
                        cell.alignment = Alignment(vertical='top', wrap_text=True, horizontal='left')
    
    def _add_borders(self, ws, df: pd.DataFrame) -> None:
        """
        添加边框
        
        Args:
            ws: 工作表对象
            df: 数据DataFrame
        """
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # 为所有单元格添加边框
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, 
                               min_col=1, max_col=len(df.columns)):
            for cell in row:
                cell.border = thin_border
    
    def _add_auto_filter(self, ws, df: pd.DataFrame) -> None:
        """
        添加自动筛选
        
        Args:
            ws: 工作表对象
            df: 数据DataFrame
        """
        # 设置自动筛选范围
        filter_range = f"A1:{get_column_letter(len(df.columns))}{ws.max_row}"
        ws.auto_filter.ref = filter_range
    
    def generate_summary_report(self, data: List[Dict[str, Any]], 
                               validation_result: Optional[Dict[str, Any]] = None) -> str:
        """
        生成汇总报告
        
        Args:
            data: 数据列表
            validation_result: 验证结果（可选）
            
        Returns:
            汇总报告文本
        """
        report_lines = []
        report_lines.append("=" * 60)
        report_lines.append("夸克校园招聘数据爬取汇总报告")
        report_lines.append("=" * 60)
        
        # 基本信息
        report_lines.append(f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        report_lines.append(f"数据总量: {len(data)} 条")
        
        # 类别统计
        category_counts = {}
        for item in data:
            category = item.get('职位类别', '未分类')
            category_counts[category] = category_counts.get(category, 0) + 1
        
        report_lines.append("\n职位类别分布:")
        for category, count in sorted(category_counts.items(), key=lambda x: x[1], reverse=True):
            report_lines.append(f"  {category}: {count} 个 ({count/len(data):.1%})")
        
        # 地点统计
        location_counts = {}
        for item in data:
            location = item.get('办公地点', '未指定')
            location_counts[location] = location_counts.get(location, 0) + 1
        
        report_lines.append("\n办公地点分布:")
        for location, count in sorted(location_counts.items(), key=lambda x: x[1], reverse=True)[:10]:
            report_lines.append(f"  {location}: {count} 个")
        
        # 学历要求统计
        education_counts = {}
        for item in data:
            education = item.get('学历要求', '未指定')
            education_counts[education] = education_counts.get(education, 0) + 1
        
        report_lines.append("\n学历要求分布:")
        for education, count in sorted(education_counts.items(), key=lambda x: x[1], reverse=True):
            report_lines.append(f"  {education}: {count} 个")
        
        # 工作年限统计
        experience_counts = {}
        for item in data:
            experience = item.get('工作年限', '未指定')
            experience_counts[experience] = experience_counts.get(experience, 0) + 1
        
        report_lines.append("\n工作年限分布:")
        for experience, count in sorted(experience_counts.items(), key=lambda x: x[1], reverse=True):
            report_lines.append(f"  {experience}: {count} 个")
        
        # 验证结果
        if validation_result:
            report_lines.append("\n数据质量验证:")
            report_lines.append(f"  有效数据: {validation_result.get('valid_positions', 0)} 条")
            report_lines.append(f"  无效数据: {validation_result.get('invalid_positions', 0)} 条")
            report_lines.append(f"  质量评分: {validation_result.get('quality_score', 0):.1%}")
            
            errors = validation_result.get('errors', [])
            if errors:
                report_lines.append(f"  错误详情: {len(errors)} 条数据有问题")
        
        report_lines.append("\n" + "=" * 60)
        
        return "\n".join(report_lines)
    
    def save_report(self, report_text: str, filename: Optional[str] = None) -> str:
        """
        保存报告文件
        
        Args:
            report_text: 报告文本
            filename: 文件名（可选）
            
        Returns:
            报告文件路径
        """
        if filename is None:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"report_{timestamp}.txt"
        
        filepath = os.path.join(self.output_dir, filename)
        
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(report_text)
        
        logger.info(f"报告已保存: {filepath}")
        return filepath